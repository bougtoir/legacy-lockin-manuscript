"""Cover letter docx, Times New Roman 12pt."""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
doc = Document()
st = doc.styles["Normal"]; st.font.name="Times New Roman"; st.font.size=Pt(12)
st.font.color.rgb=RGBColor(0,0,0)
st.element.rPr.rFonts.set(qn("w:eastAsia"),"Times New Roman")

LINES = [
"Dear Editors,",
"",
"We submit our manuscript 'When does inherited structure become adaptation lock-in?' for consideration as an Article in Nature Sustainability.",
"",
"Lock-in is one of the most widely invoked concepts in climate adaptation and sustainability transitions research. The intuition is familiar: inherited crop portfolios, irrigation networks, built environments and institutional programmes constrain future transformation. The concept carries heavy theoretical baggage — path dependence, stranded assets, carbon lock-in — and is routinely used to explain why adaptation lags. Yet it is rarely tested as a falsifiable, cross-domain prediction.",
"",
"This paper reports a coordinated empirical stress test of that prediction. Across four independent human adaptation systems — national crop specialization, subnational irrigation dependence, urban built capital under flood exposure, and the institutional initiation of flood-related property acquisition and relocation programmes — we asked whether systems carrying greater inherited specialization, fixed capital or institutional commitment show systematically attenuated transformation under stronger environmental demand. Critically, each primary design was frozen, outcome-blind, before its headline estimate was computed, and each primary test was executed exactly once.",
"",
"The results do not rescue the general lock-in hypothesis — and we did not write the paper that hypothesis wanted. Two agricultural scales produced tightly centred nulls. The one precise primary interaction, in urban built capital, is small, inland-concentrated, and reverses under prespecified respecifications. At the institutional transition margin, the legacy interaction is directional but uncertain, while the capacity-enablement leg of the mechanism is unsupported. Secondary analyses that strengthen the hypothesis are reported as such — secondary, and in need of independent confirmation — rather than silently promoted.",
"",
"We believe this is a valuable result for the readership of Nature Sustainability for three reasons. First, it converts a rhetorically powerful but loosely specified intuition into a set of explicit, examinable claims — and reports where those claims hold and where they fail. Second, it provides a defensible boundary: accumulated legacy alone is insufficient evidence of adaptation lock-in, so 'lock-in' should be diagnosed against measurable barriers rather than inferred from accumulated structure. Third, it demonstrates that null and uncertain results from frozen confirmatory designs can be reported transparently at scale, which matters for a literature in which negative findings on adaptation constraint are easy to file away.",
"",
"The manuscript is within Nature Sustainability's Article format: a 7-word title, ~150-word unreferenced abstract, ~3,400 words of main text, six display items and 43 references. All data and code are available in the accompanying reproducibility bundle and public mirrors, with acquisition metadata for third-party sources that cannot be redistributed.",
"",
"This manuscript is original, has not been published previously, and is not under consideration elsewhere. All authors have approved the submission and declare no competing interests.",
"",
"We hope you will find it suitable for peer review at Nature Sustainability.",
"",
"Sincerely,",
"",
"Tatsuki Onishi",
"on behalf of all authors",
]
for ln in LINES:
    p = doc.add_paragraph()
    r = p.add_run(ln); r.font.size=Pt(12); r.font.name="Times New Roman"
    r.font.color.rgb=RGBColor(0,0,0)
doc.save(os.path.join(ROOT,"COVER_LETTER_NS.docx"))
print("cover letter saved")
