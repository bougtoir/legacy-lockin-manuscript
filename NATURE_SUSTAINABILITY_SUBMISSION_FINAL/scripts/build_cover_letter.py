"""Cover letter docx, Times New Roman 12pt."""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
doc = Document()
st = doc.styles["Normal"]; st.font.name="Times New Roman"; st.font.size=Pt(12)
st.font.color.rgb=RGBColor(0,0,0)
st.element.rPr.rFonts.set(qn("w:eastAsia"),"Times New Roman")

LINES = [
"Dear Editors,",
"",
"I submit the manuscript 'When does inherited structure become adaptation lock-in?' for consideration as an Article in Nature Sustainability.",
"",
"Lock-in is one of the most widely invoked concepts in climate adaptation and sustainability transitions research. The intuition is familiar: inherited crop portfolios, irrigation networks, built environments and institutional programmes constrain future transformation. The concept carries heavy theoretical baggage — path dependence, stranded assets, carbon lock-in — and is routinely used to explain why adaptation lags. Yet it is rarely tested as a falsifiable, cross-domain prediction.",
"",
"This paper reports a coordinated empirical stress test of that prediction. Across four independent human adaptation systems — national crop specialization, subnational irrigation dependence, urban built capital under flood exposure, and the institutional initiation of flood-related property acquisition and relocation programmes — the programme asked whether systems carrying greater inherited specialization, fixed capital or institutional commitment show systematically attenuated transformation under stronger environmental demand. Critically, each primary design was frozen, outcome-blind, before its headline estimate was computed, and each primary test was executed exactly once.",
"",
"The complete evidence did not support the general lock-in prediction, and the paper retains the null and uncertain primary results rather than promoting stronger secondary estimates. Neither agricultural scale provided evidence of conditioning. The one precise primary interaction, in urban built capital, is small, inland-concentrated and not stable under prespecified respecifications. At the institutional transition margin the legacy interaction is negative but uncertain, and the capacity-enablement leg of the mechanism is unsupported. Secondary analyses that strengthen the hypothesis are reported as such — secondary, and in need of independent confirmation.",
"",
"I believe this is a valuable result for the readership of Nature Sustainability for three reasons. First, it converts a rhetorically powerful but loosely specified intuition into a set of explicit, examinable claims — and reports where those claims hold and where they fail. Second, it provides a defensible boundary: accumulated legacy alone is insufficient evidence of adaptation lock-in, so 'lock-in' should be diagnosed against measurable barriers rather than inferred from accumulated structure. Third, it demonstrates that null and uncertain results from frozen confirmatory designs can be reported transparently at scale, which matters for a literature in which negative findings on adaptation constraint are easy to file away.",
"",
"The manuscript is within Nature Sustainability's Article format: a 7-word title, an unreferenced abstract of fewer than 150 words, ~2,600 words of main text, five display items and 48 references. All data and code are available in the accompanying reproducibility bundle and in the public repository at https://github.com/bougtoir/legacy-lockin-manuscript, with acquisition metadata for third-party sources that cannot be redistributed.",
"",
"Throughout the programme, outcome-blind locking means that each design — estimand, sample, estimators and inference — was frozen before the corresponding headline estimate was computed; it is not external preregistration and the paper states so explicitly.",
"",
"I confirm that this manuscript is original, has not been published previously, and is not under consideration elsewhere. I declare no competing interests.",
"",
"I hope you will find it suitable for peer review at Nature Sustainability.",
"",
"Sincerely,",
"",
"Onishi Tatsuki",
]
for ln in LINES:
    p = doc.add_paragraph()
    r = p.add_run(ln); r.font.size=Pt(12); r.font.name="Times New Roman"
    r.font.color.rgb=RGBColor(0,0,0)
doc.save(os.path.join(ROOT,"COVER_LETTER_NS_FINAL.docx"))
print("cover letter saved")
