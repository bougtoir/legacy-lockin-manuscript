"""Build MAIN_MANUSCRIPT_NS_inline_FINAL.docx and _clean.docx from main_text.md."""
import re, os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MD = os.path.join(ROOT, "manuscript", "main_text.md")
FIGDIR = os.path.join(ROOT, "figures")

TABLE1 = [
 ["study","legacy","outcome","primary estimate (95% CI)","classification"],
 ["national agriculture (N=147)","crop portfolio HHI","crop-mix JSD","+0.005 [\u22120.131, +0.143], p\u22481.0","NULL"],
 ["subnational agriculture (N=340)","irrigation share 2000","crop-mix JSD","\u22120.038 [\u22120.251, +0.077], p=0.384","NULL"],
 ["urban settlements (N=4,598)","ln built volume 2000","net hazard-avoidance reallocation","\u22120.0083 [\u22120.0166, \u22120.0033], p=0.002","WEAK / CONTEXT DEPENDENT"],
 ["transition H1 legacy (N=51,686)","housing-stock age","first programme entry","\u22120.192 [\u22120.454, +0.069], p=0.149","DIRECTIONAL BUT UNCERTAIN"],
 ["transition H2 capacity (N=51,686)","prior departure obligations","first programme entry","+0.034 [\u22120.097, +0.166], p=0.610","NULL"],
 ["transition H1 lag-1 (secondary)","housing-stock age","first programme entry","\u22120.380 [\u22120.678, \u22120.082], p=0.012","SECONDARY"],
 ["transition H1 intensive (secondary)","housing-stock age","entry volume","\u22120.548 [\u22120.910, \u22120.187], p=0.003","SECONDARY"],
]

import csv as _csv
with open(os.path.join(ROOT,"extended_data","ED_Table_3.csv"),"w",newline="") as f:
    _csv.writer(f).writerows([
 ["study","unit","N","legacy","demand","outcome","estimand","estimate","95% CI","p","classification"],
 ["national agriculture","country","147","crop portfolio HHI","crop-calendar mismatch","crop-mix JSD","legacy x mismatch (fractional logit)","+0.005","[-0.131, +0.143]","~1.0","NULL"],
 ["subnational agriculture","admin-1","340","irrigation share 2000","crop-calendar mismatch","crop-mix JSD","legacy x mismatch (OLS + country FE)","-0.038","[-0.251, +0.077]","0.384","NULL"],
 ["urban settlements","urban centre","4,598","ln built volume 2000","baseline RP10 flood share","net hazard-avoidance reallocation","legacy x exposure (OLS + country FE)","-0.0083","[-0.0166, -0.0033]","0.002","WEAK / CONTEXT DEPENDENT"],
 ["transition H1 (legacy)","county-year","51,686","housing-stock age","3-yr lag flood declarations","first programme entry","demand x legacy (cloglog)","-0.192","[-0.454, +0.069]","0.149","DIRECTIONAL BUT UNCERTAIN"],
 ["transition H2 (capacity)","county-year","51,686","prior departure obligations","3-yr lag flood declarations","first programme entry","demand x capacity (cloglog)","+0.034","[-0.097, +0.166]","0.610","NULL"],
 ["transition H1 lag-1 (secondary)","county-year","51,686","housing-stock age","1-yr lag declarations","first programme entry","demand x legacy (cloglog)","-0.380","[-0.678, -0.082]","0.012","SECONDARY"],
 ["transition H1 intensive (secondary)","county-year","65,919","housing-stock age","1-yr lag declarations","entry volume","demand x legacy (PPML)","-0.548","[-0.910, -0.187]","0.003","SECONDARY"]])

def style_doc(doc):
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"; st.font.size = Pt(12)
    st.font.color.rgb = RGBColor(0,0,0)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for h in ("Heading 1","Heading 2","Heading 3"):
        s = doc.styles[h]; s.font.name="Times New Roman"; s.font.color.rgb=RGBColor(0,0,0)
        s.element.rPr.rFonts.set(qn("w:eastAsia"),"Times New Roman")

def _emit(doc_p, text, size, bold, italic):
    for seg in re.split(r"(\*\*.+?\*\*|\*.+?\*)", text):
        if not seg:
            continue
        b, it, s = bold, italic, seg
        if s.startswith("**") and s.endswith("**") and len(s) > 4:
            b, s = True, s[2:-2]
        elif s.startswith("*") and s.endswith("*") and len(s) > 2:
            it, s = True, s[1:-1]
        r = doc_p.add_run(s)
        r.bold = b; r.italic = it
        r.font.size = Pt(size); r.font.name = "Times New Roman"
        r.font.color.rgb = RGBColor(0,0,0)

def add_para(doc, text, size=12, bold=False, italic=False):
    p = doc.add_paragraph()
    parts = re.split(r"(\{[0-9,\u2013\s]+\})", text)
    for part in parts:
        m = re.fullmatch(r"\{([0-9,\u2013\s]+)\}", part)
        if m:
            r = p.add_run("[" + m.group(1).strip() + "]")
            r.font.superscript = True
            r.font.size = Pt(size); r.font.name = "Times New Roman"
            r.font.color.rgb = RGBColor(0,0,0)
        else:
            _emit(p, part, size, bold, italic)
    return p

def add_caption(doc, text):
    p = add_para(doc, text, size=10)
    p.paragraph_format.space_before = Pt(14)
    return p

def add_table1(doc):
    t = doc.add_table(rows=len(TABLE1), cols=len(TABLE1[0]))
    t.style = "Table Grid"
    for i,row in enumerate(TABLE1):
        for j,val in enumerate(row):
            cell = t.cell(i,j); cell.text = ""
            r = cell.paragraphs[0].add_run(val)
            r.font.size = Pt(9); r.font.name="Times New Roman"
            r.font.color.rgb = RGBColor(0,0,0)
            if i==0: r.bold = True

def build(inline: bool, out: str):
    doc = Document(); style_doc(doc)
    sec = doc.sections[0]
    for m in (sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin):
        pass
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    lines = open(MD, encoding="utf8").read().split("\n")
    fig_caption = {}
    for ln in lines:
        if ln.startswith("**Figure "):
            num = re.match(r"\*\*Figure (\d)", ln).group(1)
            fig_caption[num] = ln.replace("**", "").strip()
    inserted_figs = set(); ed3_done = False; skip_refs = False
    i = 0
    while i < len(lines):
        ln = lines[i]; i += 1
        if ln.startswith("## References"):
            # flush references block separately
            doc.add_heading("References", level=1)
            while i < len(lines):
                rl = lines[i]; i += 1
                if rl.startswith("## "): i -= 1; break
                if rl.strip():
                    add_para(doc, rl, size=10)
            continue
        if ln.startswith("## Figure legends"):
            doc.add_heading("Figure legends", level=1)
            while i < len(lines):
                rl = lines[i]; i += 1
                if rl.startswith("## Extended Data"): i -= 1; break
                if rl.strip():
                    add_para(doc, rl, size=10)
            continue
        if ln.startswith("## Extended Data"):
            doc.add_heading("Extended Data", level=1)
            while i < len(lines):
                rl = lines[i]; i += 1
                if rl.strip():
                    add_para(doc, rl, size=10)
                    if rl.startswith("**Extended Data Table 3") and not ed3_done:
                        ed3_done = True
                        if inline:
                            add_table1(doc)
            continue
        if ln.startswith("## "):
            doc.add_heading(ln[3:].strip(), level=1); continue
        if ln.startswith("### "):
            doc.add_heading(ln[4:].strip(), level=2); continue
        if ln.startswith("# "):
            h = doc.add_heading(ln[2:].strip(), level=0); continue
        if not ln.strip():
            continue
        add_para(doc, ln)
        # insert figures right after first citation
        if inline:
            for num in re.findall(r"Fig\.\s*(\d)", ln):
                if num not in inserted_figs and os.path.exists(f"{FIGDIR}/Figure_{num}.png"):
                    doc.add_picture(f"{FIGDIR}/Figure_{num}.png", width=Inches(6.3))
                    if num in fig_caption:
                        add_caption(doc, fig_caption[num])
                    inserted_figs.add(num)
    doc.save(out)
    print("saved", out, "figs:", sorted(inserted_figs))

build(True, os.path.join(ROOT, "MAIN_MANUSCRIPT_NS_inline_FINAL.docx"))
build(False, os.path.join(ROOT, "MAIN_MANUSCRIPT_NS_clean_FINAL.docx"))
